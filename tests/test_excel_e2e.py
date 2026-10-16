from pathlib import Path

import httpx
from openpyxl import load_workbook

from allincheck.config import Settings
from allincheck.excel import build_workbook, safe_cell
from allincheck.models import Dealer
from allincheck.rules import load_rules
from allincheck.scan import load_dealers, scan_dealer

from .conftest import FIX, make_fetcher

RULES = load_rules(Path(__file__).parent.parent / "rules" / "on.yaml")


def test_safe_cell_blocks_formulas() -> None:
    assert safe_cell('=HYPERLINK("http://evil")').startswith("'=")
    assert safe_cell("+1") == "'+1"
    assert safe_cell("normal") == "normal"
    assert len(safe_cell("x" * 5000)) == 2000


def test_end_to_end_with_saved_pages(settings: Settings, tmp_path: Path) -> None:
    pages = {
        "/": '<a href="/vehicle/1">a</a><a href="/vehicle/2">b</a><a href="/about">x</a>',
        "/vehicle/1": (FIX / "vdp_fee.html").read_text(),
        "/vehicle/2": (FIX / "vdp_ok.html").read_text(),
    }

    def handler(req: httpx.Request) -> httpx.Response:
        if req.url.path in pages:
            return httpx.Response(200, text=pages[req.url.path])
        return httpx.Response(404)

    fetcher = make_fetcher(settings, httpx.MockTransport(handler))
    dealer = Dealer(name="Test Motors", website_url="https://test.ca/", province="ON")
    result = scan_dealer(dealer, fetcher, RULES, settings)
    assert len(result.listings) == 2
    assert [f.rule_id for f in result.findings] == ["DRAFT-FEE-01"]

    out = tmp_path / "r.xlsx"
    build_workbook([result], RULES, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Findings", "To review", "Listings", "Errors", "Rules"]
    assert wb["Findings"].max_row == 2
    assert wb["Findings"].freeze_panes == "A2"


def test_blocked_dealer_marked_and_not_bypassed(settings: Settings) -> None:
    def handler(req: httpx.Request) -> httpx.Response:
        return httpx.Response(404) if req.url.path == "/robots.txt" else httpx.Response(403)

    fetcher = make_fetcher(settings, httpx.MockTransport(handler))
    dealer = Dealer(name="Blocked", website_url="https://b.ca/", province="ON")
    result = scan_dealer(dealer, fetcher, RULES, settings)
    assert result.status == "blocked"


def test_load_dealers_rejects_private(tmp_path: Path) -> None:
    csv = tmp_path / "d.csv"
    csv.write_text("name,website_url,province,notes\nA,http://127.0.0.1/,ON,\nB,notaurl,ON,\n")
    dealers, problems = load_dealers(csv)
    assert dealers == [] and len(problems) == 2


def test_audit_command_writes_summary(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    import allincheck.cli as cli
    from allincheck.fetch import PoliteFetcher as PF

    pages = {"/": '<a href="/vehicle/1">a</a>', "/vehicle/1": (FIX / "vdp_fee.html").read_text()}

    def handler(req: httpx.Request) -> httpx.Response:
        if req.url.path in pages:
            return httpx.Response(200, text=pages[req.url.path])
        return httpx.Response(404)

    def fake(settings):  # type: ignore[no-untyped-def]
        return PF(
            settings,
            transport=httpx.MockTransport(handler),
            resolver=lambda h: ["93.184.216.34"],
            sleep=lambda s: None,
        )

    monkeypatch.setattr(cli, "PoliteFetcher", fake)
    from allincheck.urlguard import check_public_url as real_check

    monkeypatch.setattr(
        cli, "check_public_url", lambda u: real_check(u, lambda h: ["93.184.216.34"])
    )
    monkeypatch.setenv("ALLINCHECK_REQUEST_DELAY_S", "0")
    cli.get_settings.cache_clear() if hasattr(cli.get_settings, "cache_clear") else None
    rc = cli.main(["audit", "demo.ca", "--out", str(tmp_path)])
    assert rc == 0
    text = (tmp_path / "demo.ca" / "summary.md").read_text()
    assert "Likely issues: 1" in text and "plus taxes and fees" in text.lower()
    assert cli.main(["audit", "http://127.0.0.1", "--out", str(tmp_path)]) == 2
